"""Import a real portfolio / watchlist from an uploaded CSV or Excel file.

The dashboard's seeded book is demo data. This module parses a user-provided CSV or
.xlsx into the app's ``Holding`` / ``WatchlistItem`` shapes so the uploaded lists become
the source of truth (stored via ``store.py`` and preferred over the seed everywhere).

Design notes:
- Column mapping is **flexible and case-insensitive** (``Symbol``/``Ticker``/``SYM`` all
  map to ``symbol``); a documented template lists every accepted column.
- A real brokerage export carries positions but NOT the technical indicators the signal
  engine needs (RSI/MACD/SMA/…). Those are **optional** columns; when a row has none,
  the holding is flagged ``has_signal_inputs = False`` so the UI shows "signals need
  market data" instead of a fabricated score — while value, weights, concentration,
  portfolio-fit and ratings still work from the position data. Never guessed.
- Exposure (``weight``) is taken from a weight column when present, else computed from
  ``quantity * price``. If neither is available the import is rejected with a clear reason.
"""
from __future__ import annotations

import csv
from datetime import datetime, timezone
import io
import re
from typing import Any

from .company_names import resolve_company_name
from .models import Holding, WatchlistItem
from .sectors import resolve_sector

# gauge/engine technical inputs — presence of any of these marks a row as signal-ready.
TECHNICAL_FIELDS = {
    "rsi", "macd_bullish", "above_sma_50", "above_sma_200",
    "relative_volume", "breakout_20d", "trend_slope_positive",
}

PORTFOLIO_ALIASES: dict[str, list[str]] = {
    "symbol": [
        "symbol", "ticker", "sym", "holding", "holdings", "asset", "assets",
        "security", "securities", "stock", "stocks", "code", "instrument",
        "position", "underlying", "item", "financial_instrument", "identifier",
        "cusip", "isin", "contract", "product", "fund", "investment", "name",
        "description", "security_description", "security_name", "holding_name",
        "company_name", "company", "symbol_description", "ticker_symbol", "sec_id", "id"
    ],
    "name": ["name", "security", "company", "description", "security_name", "title", "holding_name", "company_name", "security_description"],
    "sector": ["sector", "industry", "category", "asset_class", "group"],
    "quantity": [
        "quantity", "qty", "shares", "units", "position", "share_count", "current_shares",
        "num_shares", "holding_quantity", "holding_qty", "share_quantity", "no_of_shares",
        "number_of_shares", "total_shares", "amount", "size", "open_qty", "position_qty",
        "available_shares", "security_shares", "vol", "volume_held", "balance", "total_qty"
    ],
    "market_value": [
        "market_value", "marketvalue", "mkt_val", "mkt_value", "current_value", "value",
        "position_value", "total_value", "equity", "cur_val", "holding_value", "amount_val"
    ],
    "price": ["price", "last", "last_price", "market_price", "current_price", "close", "mark", "unit_price", "latest_price"],
    "avg_cost": ["avg_cost", "avgcost", "average_cost", "cost", "cost_basis", "costbasis", "avg_price", "average_price", "unit_cost", "purchase_price"],
    "weight": ["weight", "weight_pct", "allocation", "allocation_pct", "pct", "percent", "target_weight", "portfolio_pct", "weighting"],
    "day_change": ["day_change", "daychange", "change_pct", "day_change_pct", "todays_change", "change", "daily_change"],
    "rsi": ["rsi", "rsi_14"],
    "macd_bullish": ["macd_bullish", "macd", "macdbullish", "macd_signal"],
    "above_sma_50": ["above_sma_50", "above_sma50", "sma50", "above_50d", "above_ma50", "ma50"],
    "above_sma_200": ["above_sma_200", "above_sma200", "sma200", "above_200d", "above_ma200", "ma200"],
    "relative_volume": ["relative_volume", "rel_vol", "rvol", "relvol", "volume_ratio"],
    "breakout_20d": ["breakout_20d", "breakout", "breakout20d", "new_high_20d", "high_20d"],
    "trend_slope_positive": ["trend_slope_positive", "trend_slope", "trend_positive", "sma50_rising", "uptrend"],
}

WATCHLIST_ALIASES: dict[str, list[str]] = {
    "symbol": [
        "symbol", "ticker", "sym", "holding", "holdings", "asset", "assets",
        "security", "securities", "stock", "stocks", "code", "instrument",
        "position", "underlying", "item", "financial_instrument", "identifier",
        "cusip", "isin", "contract", "product", "fund", "investment", "name",
        "description", "security_description", "security_name", "holding_name",
        "company_name", "company", "symbol_description", "ticker_symbol", "sec_id", "id"
    ],
    "name": ["name", "security", "company", "description", "holding_name"],
    "sector": ["sector", "industry", "category"],
    "note": ["note", "notes", "comment", "comments", "thesis", "reason", "rationale"],
}

_TRUE = {"true", "t", "yes", "y", "1", "bullish", "up", "positive"}


class ImportError_(ValueError):
    """Raised when a file cannot be turned into any valid rows."""


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(header).strip().lower()).strip("_")


def _resolve_columns(headers: list[str], aliases: dict[str, list[str]]) -> dict[str, str]:
    norm_to_original: dict[str, str] = {}
    for header in headers:
        norm = _norm(header)
        if norm and norm not in norm_to_original:
            norm_to_original[norm] = header
    mapping: dict[str, str] = {}
    for field, names in aliases.items():
        # Pass 1: exact match
        for name in names:
            if name in norm_to_original:
                mapping[field] = norm_to_original[name]
                break
        # Pass 2: substring match on normalized header
        if field not in mapping:
            for norm, orig in norm_to_original.items():
                if any(name in norm for name in names):
                    mapping[field] = orig
                    break
    return mapping


def _present(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _num(value: Any) -> float | None:
    if not _present(value):
        return None
    text = re.sub(r"[,$%\s]", "", str(value)).replace("(", "-").replace(")", "")
    try:
        return float(text)
    except ValueError:
        return None


def _boolean(value: Any) -> bool:
    return str(value).strip().lower() in _TRUE if _present(value) else False


def _clean_symbol(val: Any) -> str:
    if val is None:
        return ""
    if hasattr(val, "strftime"):  # datetime / date object
        return ""
    s = str(val).strip()
    if not s:
        return ""
    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,4}", s) or re.match(r"^\d{1,2}[-/]\d{1,2}[-/]\d{2,4}", s) or re.match(r"^\d{4}[-/]", s):
        return ""
    if not re.search(r"[a-zA-Z]", s):
        return ""
    lower = s.lower()
    summary_keywords = (
        "total", "totals", "account", "summary", "subtotal", "net assets", "as of",
        "page", "disclaimer", "cash balance", "grand total", "portfolio total", "cash"
    )
    if any(lower == k or lower.startswith(k) for k in summary_keywords):
        return ""
    s = re.sub(r"^[$#@]", "", s).strip()
    if ":" in s:
        parts = s.split(":")
        last = parts[-1].strip()
        if re.match(r"^[A-Za-z][A-Za-z0-9.\-]*$", last) and len(last) <= 10:
            s = last
    if " - " in s:
        s = s.split(" - ")[0].strip()
    if "(" in s:
        match = re.search(r"\(([A-Za-z0-9.\-]+)\)", s)
        if match and re.search(r"[a-zA-Z]", match.group(1)) and len(match.group(1)) <= 10:
            s = match.group(1)
        else:
            s = s.split("(")[0].strip()
    words = s.split()
    first_word = re.sub(r"[^a-zA-Z0-9.\-]", "", words[0]) if words else ""
    if re.match(r"^[A-Za-z0-9.\-]+$", first_word) and re.search(r"[a-zA-Z]", first_word) and 1 <= len(first_word) <= 10:
        s = first_word
    cleaned = s.upper()
    return cleaned if (re.search(r"[A-Z]", cleaned) and 1 <= len(cleaned) <= 10) else ""


def _is_ticker_like(val: str) -> bool:
    cleaned = _clean_symbol(val)
    return bool(cleaned)


def _parse_raw_lines(raw_rows: list[list[str]]) -> list[dict[str, Any]]:
    if not raw_rows:
        return []

    symbol_aliases = PORTFOLIO_ALIASES["symbol"]

    # Search top 30 rows for header row containing symbol/ticker alias
    header_idx = -1
    for r_idx, row in enumerate(raw_rows[:30]):
        for cell in row:
            n = _norm(cell)
            if any(alias == n or alias in n for alias in symbol_aliases):
                header_idx = r_idx
                break
        if header_idx != -1:
            break

    if header_idx != -1:
        headers = [c.strip() for c in raw_rows[header_idx]]
        records: list[dict[str, Any]] = []
        for row in raw_rows[header_idx + 1:]:
            if not row or all(str(c).strip() == "" for c in row):
                continue
            records.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
        return records

    # Headerless fallback: check if column 0..10 looks like tickers
    ticker_col = -1
    sample_cols = min(10, max(len(r) for r in raw_rows[:5])) if raw_rows else 0
    for col_i in range(sample_cols):
        match_count = sum(1 for row in raw_rows[:20] if len(row) > col_i and _is_ticker_like(str(row[col_i])))
        if match_count >= min(2, len(raw_rows)):
            ticker_col = col_i
            break

    if ticker_col != -1:
        col_count = max(len(r) for r in raw_rows)
        headers = [f"col_{i}" for i in range(col_count)]
        headers[ticker_col] = "symbol"
        other_cols = [i for i in range(col_count) if i != ticker_col]
        if len(other_cols) == 1:
            headers[other_cols[0]] = "weight"
        elif len(other_cols) >= 2:
            headers[other_cols[0]] = "quantity"
            headers[other_cols[1]] = "price"
        records = []
        for row in raw_rows:
            if not row or all(str(c).strip() == "" for c in row):
                continue
            records.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
        return records

    # Default fallback: first row is header
    headers = [str(c).strip() for c in raw_rows[0]]
    records = []
    for row in raw_rows[1:]:
        if not row or all(str(c).strip() == "" for c in row):
            continue
        records.append({headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))})
    return records


def read_tabular(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Turn an uploaded CSV or .xlsx file into a list of header->value dicts."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            import openpyxl  # lazy: only needed for Excel uploads
        except ModuleNotFoundError as exc:  # pragma: no cover - dependency guard
            raise ImportError_("Excel support needs the 'openpyxl' package on the server; upload a CSV, or install it.") from exc
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheetnames = workbook.sheetnames
        if not sheetnames:
            workbook.close()
            return []

        # If only 1 sheet, use it directly
        if len(sheetnames) == 1:
            sheet = workbook.active
            raw_rows = list(sheet.iter_rows(values_only=True))
            workbook.close()
            if not raw_rows:
                return []
            rows_str = []
            for r in raw_rows:
                if r is None or all(cell is None or str(cell).strip() == "" for cell in r):
                    continue
                rows_str.append([str(cell).strip() if cell is not None else "" for cell in r])
            return _parse_raw_lines(rows_str)

        # Multi-sheet workbook: score all sheets to find the best portfolio / holdings sheet
        best_rows_str: list[list[str]] = []
        best_score = -999
        for sname in sheetnames:
            sheet = workbook[sname]
            raw_rows = list(sheet.iter_rows(values_only=True))
            if not raw_rows:
                continue
            rows_str = []
            for r in raw_rows:
                if r is None or all(cell is None or str(cell).strip() == "" for cell in r):
                    continue
                rows_str.append([str(cell).strip() if cell is not None else "" for cell in r])
            if not rows_str:
                continue
            parsed = _parse_raw_lines(rows_str)
            if not parsed:
                continue
            cols = _resolve_columns(list(parsed[0].keys()), PORTFOLIO_ALIASES)
            if "symbol" not in cols:
                continue
            score = 50
            if "quantity" in cols:
                score += 100
            if "market_value" in cols:
                score += 30
            if "price" in cols:
                score += 20
            if "weight" in cols:
                score += 15
            if "avg_cost" in cols:
                score += 10
            slower = sname.lower()
            if any(k in slower for k in ("holding", "position", "portfolio", "share", "stock", "equity", "excel", "allocation")):
                score += 25
            if any(k in slower for k in ("summary", "overview", "rating", "dividend", "performance", "metric", "disclaimer", "note")):
                score -= 15
            valid_tickers = sum(1 for r in parsed[:30] if _clean_symbol(r.get(cols["symbol"])))
            score += min(valid_tickers, 25)
            if score > best_score:
                best_score = score
                best_rows_str = rows_str

        workbook.close()
        if best_rows_str:
            return _parse_raw_lines(best_rows_str)

        # Fallback to active sheet
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
        rows_str = []
        for r in raw_rows:
            if r is None or all(cell is None or str(cell).strip() == "" for cell in r):
                continue
            rows_str.append([str(cell).strip() if cell is not None else "" for cell in r])
        return _parse_raw_lines(rows_str)

    # Default CSV handling: decode with fallbacks for encodings / UTF-16
    text = ""
    for enc in ("utf-8-sig", "utf-16", "utf-16le", "utf-16be", "latin1", "cp1252"):
        try:
            text = content.decode(enc)
            if text:
                break
        except Exception:
            continue
    if not text:
        text = content.decode("utf-8", errors="replace")

    text = text.replace("\0", "").strip()
    if not text:
        return []

    lines = [l.strip() for l in text.splitlines() if l.strip()]
    if not lines:
        return []

    sample = lines[0]
    delim = ","
    if "\t" in sample:
        delim = "\t"
    elif ";" in sample and "," not in sample:
        delim = ";"
    elif "|" in sample and "," not in sample:
        delim = "|"

    raw_rows = []
    for line in lines:
        f = list(csv.reader([line], delimiter=delim))[0] if line else []
        if f:
            raw_rows.append([val.strip() for val in f])

    return _parse_raw_lines(raw_rows)


def parse_portfolio(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Parse tabular rows into holdings + a signal-inputs flag + warnings."""
    if not rows:
        raise ImportError_("The file has no rows.")
    columns = _resolve_columns(list(rows[0].keys()), PORTFOLIO_ALIASES)
    if "symbol" not in columns:
        raise ImportError_("No symbol/ticker column found. Include a 'symbol' column.")
    technical_columns = TECHNICAL_FIELDS & set(columns)
    warnings: list[str] = []

    staged: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=2):
        symbol = _clean_symbol(row.get(columns["symbol"]))
        if not symbol:
            warnings.append(f"Row {index}: no symbol — skipped.")
            continue
        qty = _num(row.get(columns["quantity"])) if "quantity" in columns else None
        price = _num(row.get(columns["price"])) if "price" in columns else None
        weight = _num(row.get(columns["weight"])) if "weight" in columns else None
        avg_cost = _num(row.get(columns["avg_cost"])) if "avg_cost" in columns else None
        market_value = _num(row.get(columns["market_value"])) if "market_value" in columns else None
        day_change = _num(row.get(columns["day_change"])) if "day_change" in columns else None

        # Cross-derive missing fields among quantity, price, and market_value
        if (qty is None or qty == 0.0) and market_value is not None and price and price > 0:
            qty = round(market_value / price, 4)
        if market_value is None and qty is not None and price is not None:
            market_value = round(qty * price, 2)
        if (price is None or price == 0.0) and market_value is not None and qty and qty > 0:
            price = round(market_value / qty, 4)

        row_has_signal = any(_present(row.get(columns[f])) for f in technical_columns)
        staged.append({
            "symbol": symbol,
            "name": resolve_company_name(symbol, str(row.get(columns["name"])).strip() if "name" in columns and _present(row.get(columns["name"])) else None),
            "sector": resolve_sector(symbol, str(row.get(columns["sector"])).strip() if "sector" in columns and _present(row.get(columns["sector"])) else None),
            "quantity": qty or 0.0,
            "price": price or 0.0,
            "avg_cost": avg_cost if avg_cost is not None else (price or 0.0),
            "day_change": day_change or 0.0,
            "weight_raw": weight,
            "market_value": market_value,
            "rsi": _num(row.get(columns["rsi"])) if "rsi" in columns else None,
            "macd_bullish": _boolean(row.get(columns["macd_bullish"])) if "macd_bullish" in columns else False,
            "above_sma_50": _boolean(row.get(columns["above_sma_50"])) if "above_sma_50" in columns else False,
            "above_sma_200": _boolean(row.get(columns["above_sma_200"])) if "above_sma_200" in columns else False,
            "relative_volume": _num(row.get(columns["relative_volume"])) if "relative_volume" in columns else None,
            "breakout_20d": _boolean(row.get(columns["breakout_20d"])) if "breakout_20d" in columns else False,
            "trend_slope_positive": _boolean(row.get(columns["trend_slope_positive"])) if "trend_slope_positive" in columns else False,
            "has_signal_inputs": bool(row_has_signal),
        })

    if not staged:
        raise ImportError_("No rows had a symbol.")

    # Determine weights: explicit weight column wins; else derive from market value.
    have_weights = "weight" in columns and any(item["weight_raw"] is not None for item in staged)
    if have_weights:
        weight_sum = sum(item["weight_raw"] or 0.0 for item in staged)
        scale = 100.0 if weight_sum and weight_sum <= 1.5 else 1.0  # fractions -> percent
        for item in staged:
            item["weight"] = round((item["weight_raw"] or 0.0) * scale, 4)
            if item["weight_raw"] is None:
                warnings.append(f"{item['symbol']}: no weight given — set to 0%.")
    else:
        total_mv = sum(item["market_value"] or 0.0 for item in staged)
        if total_mv <= 0:
            raise ImportError_("Cannot compute exposure: provide a 'weight' column, or 'quantity' and 'price' (or 'market_value') so weights can be derived.")
        for item in staged:
            item["weight"] = round((item["market_value"] or 0.0) / total_mv * 100.0, 4)

    holdings: list[Holding] = []
    for item in staged:
        holdings.append(Holding(
            symbol=item["symbol"], name=item["name"], sector=item["sector"],
            quantity=item["quantity"], price=item["price"], avgCost=item["avg_cost"],
            dayChange=item["day_change"], weight=item["weight"],
            rsi=item["rsi"] if item["rsi"] is not None else 50.0,
            macdBullish=item["macd_bullish"], aboveSma50=item["above_sma_50"],
            aboveSma200=item["above_sma_200"],
            relativeVolume=item["relative_volume"] if item["relative_volume"] is not None else 0.0,
            breakout20d=item["breakout_20d"], trendSlopePositive=item["trend_slope_positive"],
            hasSignalInputs=item["has_signal_inputs"],
        ))

    return {
        "holdings": [h.model_dump(by_alias=True) for h in holdings],
        "has_signal_inputs": bool(technical_columns),
        "warnings": warnings,
        "count": len(holdings),
    }


def parse_watchlist(rows: list[dict[str, Any]]) -> dict[str, Any]:
    """Parse tabular rows into watchlist items + warnings."""
    if not rows:
        raise ImportError_("The file has no rows.")
    columns = _resolve_columns(list(rows[0].keys()), WATCHLIST_ALIASES)
    if "symbol" not in columns:
        raise ImportError_("No symbol/ticker column found. Include a 'symbol' column.")
    warnings: list[str] = []
    seen: set[str] = set()
    items: list[WatchlistItem] = []
    for index, row in enumerate(rows, start=2):
        symbol = _clean_symbol(row.get(columns["symbol"]))
        if not symbol:
            warnings.append(f"Row {index}: no symbol — skipped.")
            continue
        if symbol in seen:
            warnings.append(f"{symbol}: duplicate — kept the first.")
            continue
        seen.add(symbol)
        items.append(WatchlistItem(
            symbol=symbol,
            name=resolve_company_name(symbol, str(row.get(columns["name"])).strip() if "name" in columns and _present(row.get(columns["name"])) else None),
            sector=(str(row.get(columns["sector"])).strip() if "sector" in columns and _present(row.get(columns["sector"])) else None),
            note=(str(row.get(columns["note"])).strip() if "note" in columns and _present(row.get(columns["note"])) else None),
        ))
    if not items:
        raise ImportError_("No rows had a symbol.")
    return {"items": [i.model_dump() for i in items], "warnings": warnings, "count": len(items)}


